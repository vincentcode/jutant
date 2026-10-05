"""Ingestion end to end on PostgreSQL, with a fake embedder and a fake OCR engine."""

import shutil
from pathlib import Path

import docx
import pytest
from asgiref.sync import sync_to_async
from django.conf import settings
from django.core.management import CommandError, call_command
from PIL import Image

from apps.knowledge import selectors
from apps.knowledge.models import Document
from ingestion import setup
from ingestion.parsers import pdf as pdf_parser
from ingestion.pipeline import IngestError, Ingestor
from providers.ocr.base import OcrUnavailable

pytestmark = pytest.mark.django_db(transaction=True)
LABELS = ["public", "internal"]


class Embedder:
    def __init__(self, dim: int = settings.JUTANT_EMBED_DIM) -> None:
        self.dim = dim
        self.batches: list[int] = []

    async def embed(self, texts: list[str]) -> list[list[float]]:
        self.batches.append(len(texts))
        return [[1.0] + [0.0] * (self.dim - 1) for _ in texts]


class FakeOcr:
    def __init__(self, text: str = "Scanned circular about dormant accounts.") -> None:
        self.text = text
        self.calls = 0

    def extract_text(self, file_bytes: bytes, mime_type: str) -> str:
        self.calls += 1
        return self.text


class BrokenOcr:
    def extract_text(self, file_bytes: bytes, mime_type: str) -> str:
        raise OcrUnavailable("tesseract is not installed")


def ingestor(**kwargs) -> Ingestor:
    return Ingestor(
        model=kwargs.pop("model", Embedder()),
        labels=LABELS,
        embed_dim=settings.JUTANT_EMBED_DIM,
        **kwargs,
    )


def text_pdf(path: Path, lines: list[str]) -> Path:
    """A one-page PDF with a real text layer, written by hand."""
    ops = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(f"({line}) Tj T*" for line in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(ops)} >>\nstream\n{ops}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    path.write_bytes(out)
    return path


def scanned_pdf(path: Path) -> Path:
    """A PDF whose only page is an image, with no text layer, like a scan."""
    Image.new("RGB", (200, 100), "white").save(path, "PDF")
    return path


def word_document(path: Path) -> Path:
    document = docx.Document()
    document.add_heading("KYC Policy", level=1)
    document.add_paragraph("Introduction to customer checks.")
    document.add_heading("4.2 Joint accounts", level=2)
    document.add_paragraph("Both holders must provide identification.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Document", "Required"
    table.cell(1, 0).text, table.cell(1, 1).text = "Passport", "Yes"
    document.save(path)
    return path


async def search(query: str, labels=LABELS):
    return await sync_to_async(selectors.search)(
        query=query, query_embedding=None, classifications=labels, limit=5
    )


# --- parsers ------------------------------------------------------------------


def test_pdf_text_layer_is_read(tmp_path: Path) -> None:
    parsed = pdf_parser.parse(text_pdf(tmp_path / "a.pdf", ["Fees change in May.", "Second line."]))
    assert "Fees change in May." in parsed.text
    assert (parsed.pages, parsed.ocr_pages, parsed.unreadable_pages) == (1, 0, 0)


def test_scanned_pdf_pages_go_through_ocr(tmp_path: Path) -> None:
    ocr = FakeOcr()
    parsed = pdf_parser.parse(scanned_pdf(tmp_path / "scan.pdf"), ocr)
    assert parsed.text == "Scanned circular about dormant accounts."
    assert (parsed.ocr_pages, ocr.calls) == (1, 1)


def test_scanned_pdf_without_ocr_is_unreadable(tmp_path: Path) -> None:
    parsed = pdf_parser.parse(scanned_pdf(tmp_path / "scan.pdf"), BrokenOcr())
    assert (parsed.text, parsed.unreadable_pages) == ("", 1)


# --- pipeline -----------------------------------------------------------------


async def test_word_document_is_indexed_with_its_sections(tmp_path: Path) -> None:
    result = await ingestor().ingest(
        word_document(tmp_path / "kyc_policy.docx"), doc_type="policy", classification="internal"
    )
    assert (result.status, result.chunks) == ("indexed", 2)  # one chunk per heading

    hits = await search("joint holders identification")
    assert [(c.document.title, c.section) for c, _ in hits] == [
        ("Kyc policy", "4.2 Joint accounts")
    ]
    document = await sync_to_async(Document.objects.get)(id=result.document_id)
    assert "Passport | Yes" in document.text
    assert document.version == "1"


async def test_unchanged_file_is_skipped_and_changed_file_replaces(tmp_path: Path) -> None:
    path = tmp_path / "fees.md"
    path.write_text("# Fees\n\nMonthly fee is 5 cedis.")
    first = await ingestor().ingest(path, doc_type="policy", classification="public")
    again = await ingestor().ingest(path, doc_type="policy", classification="public")
    path.write_text("# Fees\n\nMonthly fee is 7 cedis.")
    updated = await ingestor().ingest(path, doc_type="policy", classification="public")

    assert (first.status, again.status, updated.status) == ("indexed", "skipped", "replaced")
    assert updated.document_id == first.document_id
    assert await sync_to_async(Document.objects.count)() == 1
    assert [c.text for c, _ in await search("monthly fee")] == ["Monthly fee is 7 cedis."]
    document = await sync_to_async(Document.objects.get)(id=first.document_id)
    assert document.version == "2"


async def test_scanned_pdf_is_indexed_through_ocr(tmp_path: Path) -> None:
    result = await ingestor(ocr=FakeOcr()).ingest(
        scanned_pdf(tmp_path / "circular.pdf"), doc_type="circular", classification="internal"
    )
    assert result.status == "indexed"
    assert result.note == "1 of 1 pages read by OCR"


async def test_scanned_pdf_without_ocr_fails_with_a_reason(tmp_path: Path) -> None:
    with pytest.raises(IngestError, match="no text found .*OCR is not available"):
        await ingestor().ingest(
            scanned_pdf(tmp_path / "scan.pdf"), doc_type="circular", classification="internal"
        )
    assert await sync_to_async(Document.objects.count)() == 0


async def test_embeddings_are_requested_in_batches(tmp_path: Path) -> None:
    path = tmp_path / "long.md"
    path.write_text("\n\n".join(f"# Section {i}\n\nText for section {i}." for i in range(5)))
    model = Embedder()
    await ingestor(model=model, batch_size=2).ingest(
        path, doc_type="policy", classification="public"
    )
    assert model.batches == [2, 2, 1]


async def test_bad_inputs_are_refused_with_a_reason(tmp_path: Path) -> None:
    path = tmp_path / "a.md"
    path.write_text("Some text.")
    with pytest.raises(IngestError, match="not a label of this pack"):
        await ingestor().ingest(path, doc_type="policy", classification="secret")
    spreadsheet = tmp_path / "a.xls"
    spreadsheet.write_bytes(b"x")
    with pytest.raises(IngestError, match="unsupported file type"):
        await ingestor().ingest(spreadsheet, doc_type="policy", classification="public")
    with pytest.raises(IngestError, match="returns 3 dimensions"):
        await ingestor(model=Embedder(dim=3)).ingest(
            path, doc_type="policy", classification="public"
        )


async def test_reindex_rebuilds_chunks_from_stored_text(tmp_path: Path) -> None:
    path = tmp_path / "fees.md"
    path.write_text("# Fees\n\nMonthly fee is 5 cedis.")
    await ingestor().ingest(path, doc_type="policy", classification="public")
    model = Embedder()
    assert await ingestor(model=model).reindex() == 1
    assert model.batches == [1]
    assert len(await search("monthly fee")) == 1


# --- commands -----------------------------------------------------------------


def test_ingest_command_reports_each_file(tmp_path: Path, monkeypatch, capsys) -> None:
    monkeypatch.setattr(setup, "build_ingestor", lambda: ingestor(ocr=None))
    (tmp_path / "fees.md").write_text("# Fees\n\nMonthly fee is 5 cedis.")
    word_document(tmp_path / "kyc.docx")
    scanned_pdf(tmp_path / "scan.pdf")
    (tmp_path / "notes.xls").write_bytes(b"ignored: unsupported type")

    with pytest.raises(CommandError, match="1 file"):
        call_command(
            "ingest_documents",
            str(tmp_path),
            "--doc-type",
            "policy",
            "--classification",
            "internal",
        )

    out, err = capsys.readouterr()
    assert "indexed   fees.md, 1 chunks" in out
    assert "indexed   kyc.docx, 2 chunks" in out
    assert "failed    scan.pdf: scan.pdf: no text found" in err
    assert "2 of 3 files ingested." in out


def test_reindex_command(tmp_path: Path, monkeypatch, capsys) -> None:
    monkeypatch.setattr(setup, "build_ingestor", lambda: ingestor())
    (tmp_path / "fees.md").write_text("# Fees\n\nMonthly fee is 5 cedis.")
    call_command(
        "ingest_documents", str(tmp_path), "--doc-type", "policy", "--classification", "public"
    )
    call_command("reindex")
    assert "1 document(s) re-indexed." in capsys.readouterr().out


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="tesseract is not installed")
def test_real_tesseract_reads_an_image() -> None:
    from PIL import ImageDraw

    from providers.ocr.tesseract import TesseractOcr

    image = Image.new("RGB", (600, 120), "white")
    ImageDraw.Draw(image).text((20, 40), "DORMANT ACCOUNT", fill="black", font_size=40)
    import io

    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    assert "DORMANT" in TesseractOcr().extract_text(buffer.getvalue(), "image/png").upper()
